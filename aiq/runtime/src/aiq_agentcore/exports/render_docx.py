# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""DOCX export via python-docx (ADR-26): real headings, lists, tables, images, superscript citation links to
bookmarked Sources entries, external hyperlinks, footer with the package id.

python-docx 1.2 exposes no footnote API, so citations are endnote-style: ``[N]`` is a superscript
``w:hyperlink w:anchor="src_N"`` pointing at a ``w:bookmarkStart`` on the Sources entry (decision item 5,
export-rendering.md §6). pandoc/LibreOffice were rejected: no binaries in the image. Port of proto/render_docx.py.
"""

from __future__ import annotations

import io
import itertools

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .mdmodel import Block, Inline, Package

LINK_COLOR = "0F766E"  # tokens.css --accent (light)
MUTED = RGBColor(0x47, 0x55, 0x69)
OK_GREEN = RGBColor(0x15, 0x80, 0x3D)


def _run_props(color: str | None = None, underline: bool = False, superscript: bool = False, mono: bool = False):
    rpr = OxmlElement("w:rPr")
    if mono:
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), "Consolas")
        rpr.append(f)
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    if superscript:
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
    return rpr


def _text_run(text: str, **props):
    r = OxmlElement("w:r")
    r.append(_run_props(**props))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_internal_link(paragraph, text: str, anchor: str, superscript: bool = True) -> None:
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    h.set(qn("w:history"), "1")
    h.append(_text_run(text, color=LINK_COLOR, superscript=superscript))
    paragraph._p.append(h)


def add_external_link(paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    h.set(qn("w:history"), "1")
    h.append(_text_run(text, color=LINK_COLOR, underline=True))
    paragraph._p.append(h)


def add_bookmark(paragraph, name: str, ids: itertools.count) -> None:
    bid = str(next(ids))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    ppr = paragraph._p.pPr
    if ppr is not None:
        ppr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_inlines(paragraph, inlines: list[Inline], pkg: Package) -> None:
    for i in inlines:
        if i.kind == "text":
            r = paragraph.add_run(i.text)
            r.bold, r.italic, r.font.strike = i.bold or None, i.italic or None, i.strike or None
        elif i.kind == "code":
            r = paragraph.add_run(i.text)
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif i.kind == "link":
            add_external_link(paragraph, i.text, i.href or "")
        elif i.kind == "cite":
            add_internal_link(paragraph, f"[{i.n}]", f"src_{i.n}")
        elif i.kind == "br":
            paragraph.add_run().add_break(WD_BREAK.LINE)


def _picture(doc, data: bytes | None, width_in: float = 6.0) -> bool:
    if not data:
        return False
    try:
        doc.add_picture(io.BytesIO(data), width=Inches(width_in))
    except Exception:  # noqa: BLE001 — undecodable image: skip rather than fail the export
        return False
    return True


def add_blocks(doc, blocks: list[Block], pkg: Package, list_depth: int = 0) -> None:
    for b in blocks:
        if b.kind == "heading":
            p = doc.add_heading(level=min(max(b.level - 1, 1), 9))
            add_inlines(p, b.inlines, pkg)
        elif b.kind == "paragraph":
            p = doc.add_paragraph()
            add_inlines(p, b.inlines, pkg)
        elif b.kind == "list":
            base = "List Number" if b.ordered else "List Bullet"
            style = base if list_depth == 0 else f"{base} {min(list_depth + 1, 3)}"
            for item in b.items:
                first = True
                for child in item:
                    if child.kind == "paragraph" and first:
                        p = doc.add_paragraph(style=style)
                        add_inlines(p, child.inlines, pkg)
                        first = False
                    elif child.kind == "list":
                        add_blocks(doc, [child], pkg, list_depth + 1)
                    else:
                        add_blocks(doc, [child], pkg, list_depth)
        elif b.kind == "table":
            hdr = b.header[0] if b.header else []
            ncols = max([len(hdr)] + [len(r) for r in b.rows]) or 1
            table = doc.add_table(rows=(1 if hdr else 0) + len(b.rows), cols=ncols)
            table.style = "Table Grid"
            ri = 0
            if hdr:
                for ci, cell in enumerate(hdr[:ncols]):
                    p = table.rows[0].cells[ci].paragraphs[0]
                    add_inlines(p, cell, pkg)
                    for r in p.runs:
                        r.bold = True
                ri = 1
            for r in b.rows:
                for ci, cell in enumerate(r[:ncols]):
                    add_inlines(table.rows[ri].cells[ci].paragraphs[0], cell, pkg)
                ri += 1
            doc.add_paragraph()
        elif b.kind == "code":
            p = doc.add_paragraph(style="No Spacing")
            r = p.add_run(b.text)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(8)
        elif b.kind == "image":
            got = pkg.image(b.src)
            if _picture(doc, got[0] if got else None) and b.alt:
                doc.add_paragraph(b.alt, style="Caption")
        elif b.kind == "hr":
            doc.add_paragraph("—" * 20).alignment = 1
        elif b.kind == "blockquote":
            for child in b.children:
                if child.kind == "paragraph":
                    p = doc.add_paragraph(style="Quote")
                    add_inlines(p, child.inlines, pkg)
                else:
                    add_blocks(doc, [child], pkg, list_depth)


def render_docx(pkg: Package) -> tuple[bytes, dict]:
    """Render the package to DOCX bytes. Returns (docx_bytes, meta)."""
    doc = Document()
    ids = itertools.count(100)
    doc.core_properties.title = pkg.title
    doc.core_properties.subject = f"Research package {pkg.job_id}"
    doc.core_properties.author = "AI-Q on Amazon Bedrock AgentCore"
    doc.core_properties.comments = ("Every [N] citation links to the numbered Sources list; every source was retrieved "
                                    "during the job.")
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    doc.add_heading(pkg.title, 0)
    meta = f"Research package {pkg.job_id}"
    if pkg.created_at:
        meta += f" · researched {pkg.created_at[:10]}"
    meta += f" · {len(pkg.sources)} sources ({pkg.verified_count} verified)"
    mp = doc.add_paragraph(meta)
    mp.runs[0].font.size = Pt(9)
    mp.runs[0].font.color.rgb = MUTED
    add_blocks(doc, pkg.body, pkg)
    trailing = pkg.trailing_artifacts()
    pictures = 0
    if trailing:
        doc.add_heading("Artifacts", 1)
        for a in trailing:
            if a.is_image and _picture(doc, a.data):
                doc.add_paragraph(a.caption, style="Caption")
                pictures += 1
            elif not a.is_image:
                doc.add_paragraph(f"{a.caption or a.filename} ({a.mime}, {len(a.data):,} bytes)", style="Caption")
    doc.add_heading("Sources", 1)
    for s in pkg.sources:
        p = doc.add_paragraph()
        add_bookmark(p, f"src_{s.n}", ids)
        p.add_run(f"[{s.n}] {s.display_title()} — ")
        if s.url:
            add_external_link(p, s.url, s.url)
        else:
            p.add_run(s.document_key or "")
        if s.verified:
            r = p.add_run("  ✓ verified")
            r.font.size = Pt(8)
            r.font.color.rgb = OK_GREEN
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"Research package {pkg.job_id} · AI-Q on Amazon Bedrock AgentCore"
    footer.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"backend": "python-docx", "sources": len(pkg.sources), "artifact_pictures": pictures,
                            "footnotes_api": hasattr(doc, "footnotes")}
